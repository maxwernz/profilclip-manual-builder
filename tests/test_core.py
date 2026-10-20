from pathlib import Path

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from manual_builder import ooxml
from manual_builder.build import compose
from manual_builder.importer import import_file
from manual_builder.library import Library, Machine
from manual_builder.normalize import normalize_document
from manual_builder.seed import DEFAULT_LIBRARY
from manual_builder.template import create_master_template

LEGACY = Path(__file__).resolve().parents[2] / "Bedienungsanleitungen"


# ---------------------------------------------------------------- helpers

def _para_with_runs(doc, *runs, style=None):
    p = doc.add_paragraph(style=style)
    for r in runs:
        p.add_run(r)
    return p


def _template_source(path: Path) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "Originalbetriebsanleitung Schlauchbeutel-Abfüllautomat SKT H10"
    doc.sections[0].footer.paragraphs[0].text = "Gedruckt"
    doc.add_paragraph("x", style="List Number")   # makes sure numbering.xml exists
    doc.save(str(path))
    return path


def _module_doc(path: Path, title_number: str, title: str, subs: list[str], extra=None) -> Path:
    doc = Document()
    doc.add_paragraph(f"{title_number} {title}", style="Heading 2")
    doc.add_paragraph("Einleitung für {{Maschinentyp}}.")
    for i, s in enumerate(subs, 1):
        _para_with_runs(doc, f"{title_number}.", f"{i} ", s, style="Heading 2")
        doc.add_paragraph(f"Text zu {s}.")
    if extra:
        extra(doc)
    doc.save(str(path))
    return path


@pytest.fixture
def lib(tmp_path) -> Library:
    root = tmp_path / "library"
    root.mkdir()
    (root / "library.yaml").write_text(yaml.safe_dump(DEFAULT_LIBRARY, allow_unicode=True), encoding="utf-8")
    create_master_template(_template_source(tmp_path / "src.docx"), root / "templates" / "master.docx")
    (root / "modules").mkdir()
    (root / "machines").mkdir()
    lib = Library(root)

    gen = lib.new_module("allgemein", {"de": "Allgemeines"}, "general", "chapters")
    import_file(lib, _module_doc(tmp_path / "a.docx", "1", "Sicherheit", ["Gefahren"]), gen, "de")

    clip = lib.new_module("clipper", {"de": "Clipper X"}, "clipper")
    import_file(lib, _module_doc(tmp_path / "c.docx", "9.5", "Clipper X", ["Draht", "Messer"]), clip, "de")
    clip.data["faults"] = [{"code": 7, "message": {"de": "Drahtvorschub"}, "cause": {"de": "Draht verbogen"},
                            "remedy": {"de": "Draht kürzen"}}]
    clip.data["maintenance"] = [{"interval_days": 90, "interval": {"de": "3 Monate"},
                                 "task": {"de": "Fetten"}, "parts": {"de": "Schmierpunkte Clipper"}}]
    clip.save_data()

    def service_extra(doc):
        doc.add_paragraph("{{GESAMT_WARTUNGSPLAN}}")
        doc.add_paragraph("{{GESAMT_FEHLERCODES}}")
        doc.add_paragraph("{{ABSCHNITT_STOERUNGEN}}")
    svc = lib.new_module("service", {"de": "Wartung"}, "service", "chapters")
    svc.component = {"de": "Grundmaschine"}
    svc.save_meta()
    import_file(lib, _module_doc(tmp_path / "s.docx", "11", "Wartung", [], service_extra), svc, "de")
    svc.data["faults"] = [{"code": 10, "message": {"de": "Luftdruck"}, "cause": {"de": "Druckluft fehlt"},
                           "remedy": {"de": "Druck prüfen"}}]
    svc.save_data()

    lib.new_module("option-waage", {"de": "Option Waage"}, "options")
    w = lib.modules["option-waage"]
    w.data["faults"] = [{"code": 27, "message": {"de": "Waage"}, "cause": {"de": "x"}, "remedy": {"de": "y"}}]
    w.save_data()
    lib.reload_modules()
    return lib


def _headings(docx: Path):
    doc = Document(str(docx))
    out = []
    for p in doc.paragraphs:
        if p.style.name.lower().startswith("heading"):
            out.append((int(p.style.name.split()[-1]), p.text))
    return out, doc


# ---------------------------------------------------------------- low level

def test_strip_prefix_over_several_runs():
    doc = Document()
    p = _para_with_runs(doc, "9.", "5.1", " Draht", "rollenwechsel:")
    assert ooxml.strip_prefix(p._p).strip() == "9.5.1"
    ooxml.strip_suffix(p._p)
    assert p.text == "Drahtrollenwechsel"


def test_placeholder_split_over_runs():
    doc = Document()
    p = _para_with_runs(doc, "Typ {{Masch", "inen", "typ}} fertig")
    ooxml.replace_in_paragraph(p._p, {"Maschinentyp": "SKT H10"})
    assert p.text == "Typ SKT H10 fertig"


def test_normalize_levels_and_title():
    doc = Document()
    doc.add_paragraph("Vorwort ohne Überschrift")
    doc.add_paragraph("9.3 Folienwechsel", style="Heading 2")
    doc.add_paragraph("9.3.2 Abzug", style="Heading 2")
    doc.add_paragraph("1. Schritt eins")   # list step, not a heading
    doc.add_paragraph("9.4 Siegelung", style="Heading 2")
    rep = normalize_document(doc, kind="component", title="Folie")
    assert [(h.level, h.text) for h in rep.headings] == [
        (1, "Folie"), (2, "Folienwechsel"), (3, "Abzug"), (2, "Siegelung")]


def test_normalize_is_idempotent(lib):
    doc = Document(str(lib.modules["clipper"].docx_path("de")))
    rep = normalize_document(doc, kind="component", title="Clipper X")
    assert [(h.level, h.text) for h in rep.headings] == [(1, "Clipper X"), (2, "Draht"), (2, "Messer")]
    assert not any("ergänzt" in i.message for i in rep.issues)


# ---------------------------------------------------------------- library / build

def test_validation(lib):
    m = Machine(name="t", modules=["clipper"])
    msgs = [i.message for i in lib.validate(m)]
    assert any("Allgemeines" in x and "Pflicht" in x for x in msgs)
    lib.new_module("clipper2", {"de": "Clipper Y"}, "clipper")
    m = Machine(name="t", modules=["allgemein", "clipper", "clipper2", "service"])
    assert any("Nur 1×" in i.message for i in lib.validate(m))


def test_build_compose(lib, tmp_path):
    machine = Machine(name="Test", modules=["allgemein", "clipper", "option-waage", "service"],
                      variables={"Maschinentyp": "SKT H99", "Kunde": "ACME"})
    out = tmp_path / "out.docx"
    issues = compose(lib, machine, out, log=lambda *_: None)
    assert not [i for i in issues if i.level == "error"]
    heads, doc = _headings(out)
    texts = [t for _, t in heads]
    # numbers are gone, Word numbers chapters automatically
    assert "Sicherheit" in texts and "Clipper X" in texts and "Option Waage" in texts
    assert not any(t[:1].isdigit() for t in texts)
    # each component chapter gets its own data sections
    assert (2, "Wartung") in heads and (2, "Störungen") in heads
    xml = ooxml.tostring(doc.element.body)
    assert "{{" not in xml, "all placeholders and markers must be replaced"
    assert "SKT H99" in xml and "ACME" in xml
    assert "TOC \\o" in xml and "PAGEREF mb_flt_clipper" in xml
    # overview lists codes of all modules of this machine, sorted
    all_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Luftdruck" in all_text and "Drahtvorschub" in all_text and "Waage" in all_text
    assert doc.settings.element.find(qn("w:updateFields")) is not None


def test_build_other_language_falls_back(lib, tmp_path):
    machine = Machine(name="Test", language="en", modules=["allgemein", "clipper", "service"])
    issues = compose(lib, machine, tmp_path / "en.docx", log=lambda *_: None)
    assert any("keinen Text auf 'en'" in i.message for i in issues)
    xml = ooxml.tostring(Document(str(tmp_path / "en.docx")).element.body)
    assert "Remedy" in xml and "Translation of the original operating instructions" in xml


def test_reimport_archives_old_version(lib, tmp_path):
    clip = lib.modules["clipper"]
    import_file(lib, _module_doc(tmp_path / "c2.docx", "9.5", "Clipper X", ["Neu"]), clip, "de", note="neu")
    assert clip.version == 2
    assert list((clip.path / "archive").glob("v1_de_*.docx"))


@pytest.mark.skipif(not LEGACY.exists(), reason="legacy manuals not available")
def test_seed_from_legacy(tmp_path):
    from manual_builder.seed import seed_from_legacy
    lib = seed_from_legacy(LEGACY, tmp_path / "lib", log=lambda *_: None)
    assert len(lib.modules["clipper-dcx5s"].entries("faults")) > 5
    machine = Machine.load(lib.machines()[0])
    issues = compose(lib, machine, tmp_path / "x.docx", log=lambda *_: None)
    assert not [i for i in issues if i.level == "error"]
