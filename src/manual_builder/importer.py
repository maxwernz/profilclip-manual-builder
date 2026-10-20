"""Add new or updated Word files to the library."""

from __future__ import annotations

import datetime as _dt
import shutil
from pathlib import Path

from docx import Document

from .library import Library, Module
from .normalize import NormalizeReport, normalize_document


def is_word_file(path: Path) -> bool:
    return path.suffix.lower() == ".docx" and not path.name.startswith(("~$", "~", "."))


def known_placeholders(lib: Library) -> set[str]:
    return {v["key"] for v in lib.variables} | {"Ausgabe"}


def check(lib: Library, src: Path, *, kind: str = "component", title: str = "") -> NormalizeReport:
    """Dry run: what would the import do / which problems are there?"""
    if not is_word_file(src):
        rep = NormalizeReport()
        rep.add("error", f"'{src.name}' ist keine Word-Datei (.docx) oder eine Sperrdatei")
        return rep
    doc = Document(str(src))
    return normalize_document(doc, kind=kind, title=title,
                              watch_terms=list(lib.config.get("watch_terms") or []),
                              known_placeholders=known_placeholders(lib))


def import_file(lib: Library, src: Path, module: Module, lang: str, note: str = "") -> NormalizeReport:
    """Normalise src and store it as <module>/<lang>.docx (old file goes to archive/)."""
    if not is_word_file(src):
        raise ValueError(f"'{src.name}' ist keine Word-Datei (.docx)")
    doc = Document(str(src))
    rep = normalize_document(doc, kind=module.kind, title=module.label(lang),
                             watch_terms=list(lib.config.get("watch_terms") or []),
                             known_placeholders=known_placeholders(lib))
    if rep.errors:
        return rep
    target = module.docx_path(lang)
    archive = module.path / "archive"
    archive.mkdir(exist_ok=True)
    stamp = _dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    replacing = target.exists()
    if replacing:
        shutil.copyfile(target, archive / f"v{module.version}_{lang}_{stamp}.docx")
        module.version += 1
    shutil.copyfile(src, archive / f"original_v{module.version}_{lang}_{stamp}.docx")
    doc.save(str(target))
    module.add_history(note or (f"{lang}.docx aktualisiert" if replacing else f"{lang}.docx hinzugefügt"),
                       language=lang)
    module.save_meta()
    return rep


def machines_using(lib: Library, module_id: str) -> list[str]:
    from .library import Machine
    out = []
    for p in lib.machines():
        m = Machine.load(p)
        if module_id in m.modules:
            out.append(m.name)
    return out
