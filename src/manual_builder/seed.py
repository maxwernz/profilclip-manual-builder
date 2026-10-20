"""Create a new library, and import the existing (legacy) manuals once.

The legacy import also splits the big maintenance / spare-part / fault tables of
"Wartung und Störungen.docx" into the modules they belong to. The assignment is
a first draft (see ASSIGN_* below) and is written to SEED_REPORT.md for review.
"""

from __future__ import annotations

import re
import shutil
import tempfile
from pathlib import Path

import yaml
from docx import Document
from docx.oxml.ns import qn

from . import ooxml
from .library import Library, Machine, slugify
from .normalize import _heading_style_id
from .template import create_master_template

DEFAULT_LIBRARY = {
    "name": "Profil-Clip Betriebsanleitungen",
    "languages": ["de", "en", "it"],
    "default_language": "de",
    "slots": [
        {"id": "general", "title": {"de": "Allgemeines", "en": "General", "it": "Generale"},
         "max": 1, "required": True},
        {"id": "dosing", "title": {"de": "Dosierung", "en": "Dosing", "it": "Dosaggio"}, "max": 1},
        {"id": "film", "title": {"de": "Folie & Siegelung", "en": "Film & sealing",
                                 "it": "Film e saldatura"}, "max": 1},
        {"id": "clipper", "title": {"de": "Clipper", "en": "Clipper", "it": "Clipper"}, "max": 1},
        {"id": "control", "title": {"de": "Bedienung", "en": "Operation", "it": "Comando"}, "max": 1},
        {"id": "options", "title": {"de": "Optionen", "en": "Options", "it": "Opzioni"}, "max": None},
        {"id": "service", "title": {"de": "Wartung & Störungen", "en": "Maintenance & faults",
                                    "it": "Manutenzione e guasti"}, "max": 1, "required": True},
        {"id": "consumables", "title": {"de": "Verbrauchsmaterial", "en": "Consumables",
                                        "it": "Materiali di consumo"}, "max": 1},
        {"id": "ce", "title": {"de": "CE-Erklärung", "en": "CE declaration",
                               "it": "Dichiarazione CE"}, "max": 1},
    ],
    "variables": [
        {"key": "Dokumentart", "label": "Dokumentart (Kopfzeile, Titelseite)",
         "default": {"de": "Originalbetriebsanleitung",
                     "en": "Translation of the original operating instructions",
                     "it": "Traduzione delle istruzioni originali"}},
        {"key": "Maschinenbezeichnung", "label": "Maschinenbezeichnung",
         "default": {"de": "Schlauchbeutel-Abfüllautomat", "en": "Tubular bag filling machine",
                     "it": "Insaccatrice automatica"}},
        {"key": "Maschinentyp", "label": "Maschinentyp", "default": "SKT H10"},
        {"key": "Dokumentnummer", "label": "Dokumentnummer", "default": "03810000"},
        {"key": "Kunde", "label": "Kunde", "default": ""},
        {"key": "Seriennummer", "label": "Seriennummer", "default": ""},
        {"key": "Baujahr", "label": "Baujahr", "default": ""},
        {"key": "Ausgabe", "label": "Ausgabedatum (leer = heute)", "default": ""},
        # manufacturer (title page) – from www.profil-clip.de
        {"key": "Hersteller", "label": "Hersteller", "default": "Profi L-Clip GmbH"},
        {"key": "Herstelleranschrift", "label": "Anschrift",
         "default": {"de": "An der Gumme 3, D-79348 Freiamt, Deutschland",
                     "en": "An der Gumme 3, D-79348 Freiamt, Germany",
                     "it": "An der Gumme 3, D-79348 Freiamt, Germania"}},
        {"key": "Telefon", "label": "Telefon", "default": "+49 (7645) 916499"},
        {"key": "Fax", "label": "Fax", "default": "+49 (7645) 916536"},
        {"key": "E-Mail", "label": "E-Mail", "default": "info@profil-clip.de"},
        {"key": "Web", "label": "Web", "default": "www.profil-clip.de"},
    ],
    # suggestions for "Maschinentyp" (product range on www.profil-clip.de)
    "machine_types": ["SKT H10", "Spouted Pouch SP1", "Stick Pack SKT V", "Stick Pack SP40"],
    # text that should usually be a variable – reported when importing
    "watch_terms": ["SKT H10", "SKT-H10", "SKT-10"],
    "labels": {
        "de": {"printed": "Gedruckt", "contents": "Inhalt",
               "toc_hint": "Inhaltsverzeichnis wird beim Erstellen aktualisiert",
               "machine_data": "Maschinendaten", "customer": "Kunde", "serial": "Seriennummer",
               "year": "Baujahr", "document_no": "Dokumentnummer", "issue": "Ausgabe",
               "maintenance": "Wartung", "spare_parts": "Verschleiß- und Ersatzteile",
               "faults": "Störungen", "interval": "Intervall", "task": "Tätigkeit",
               "parts": "Teile / Ort", "part_no": "Teilenummer", "description": "Bezeichnung",
               "code": "Störanzeige", "message": "Meldung", "cause": "Ursache",
               "remedy": "Behebung", "module": "Baugruppe", "page": "Seite", "no": "Nr.", "manufacturer": "Hersteller"},
        "en": {"printed": "Printed", "contents": "Contents",
               "toc_hint": "Table of contents is updated when the manual is built",
               "machine_data": "Machine data", "customer": "Customer", "serial": "Serial number",
               "year": "Year of manufacture", "document_no": "Document number", "issue": "Issue",
               "maintenance": "Maintenance", "spare_parts": "Wear and spare parts",
               "faults": "Faults", "interval": "Interval", "task": "Task",
               "parts": "Parts / location", "part_no": "Part number", "description": "Description",
               "code": "Fault message", "message": "Message", "cause": "Cause",
               "remedy": "Remedy", "module": "Assembly", "page": "Page", "no": "No.", "manufacturer": "Manufacturer"},
        "it": {"printed": "Stampato", "contents": "Indice",
               "toc_hint": "L'indice viene aggiornato durante la creazione",
               "machine_data": "Dati macchina", "customer": "Cliente", "serial": "Numero di serie",
               "year": "Anno di costruzione", "document_no": "Numero documento", "issue": "Edizione",
               "maintenance": "Manutenzione", "spare_parts": "Parti di usura e ricambi",
               "faults": "Guasti", "interval": "Intervallo", "task": "Attività",
               "parts": "Parti / posizione", "part_no": "Codice", "description": "Descrizione",
               "code": "Messaggio di guasto", "message": "Messaggio", "cause": "Causa",
               "remedy": "Rimedio", "module": "Gruppo", "page": "Pagina", "no": "N.", "manufacturer": "Fabbricante"},
    },
}


def create_library(root: Path, template_source: Path) -> Library:
    root.mkdir(parents=True, exist_ok=True)
    with (root / "library.yaml").open("w", encoding="utf-8") as fh:
        yaml.safe_dump(DEFAULT_LIBRARY, fh, allow_unicode=True, sort_keys=False, width=100)
    create_master_template(template_source, root / "templates" / "master.docx")
    (root / "modules").mkdir(exist_ok=True)
    (root / "machines").mkdir(exist_ok=True)
    return Library(root)


# ---------------------------------------------------------------- legacy import

# module id -> (folder/file relative to the legacy root per language, title, slot, kind, component)
LEGACY_MODULES = [
    ("allgemein", {"de": "Deutsch/Allgemeine Anleitung/Allgemein_Deutsch_Muss angepasst werden.docx"},
     {"de": "Allgemeines"}, "general", "chapters", {"de": "Grundmaschine"}),
    ("dosierung-servo-2k", {"de": "Deutsch/Dosierung/Servo 2k Dosierung (Fischer).docx"},
     {"de": "2K-Servodosierung"}, "dosing", "component", {}),
    ("dosierung-elektromechanisch", {"de": "Deutsch/Dosierung/Dosierung Elektromechanisch.docx"},
     {"de": "Dosierung elektromechanisch"}, "dosing", "component", {}),
    ("folie-siegelung-2k", {"de": "Deutsch/Folie/Folie und Siegelung 2K (Fischer).docx"},
     {"de": "Folie und Siegelung"}, "film", "component", {}),
    ("clipper-dcx5s", {"de": "Deutsch/Clipper/Draht-Deutsch.docx"},
     {"de": "Clipper DCX5S"}, "clipper", "component", {}),
    ("clipper-dcx5fc", {"de": "Deutsch/Clipper/FC-Detusch.docx", "en": "Englisch/Clipper/FC-Englisch.docx"},
     {"de": "Clipper DCX5FC", "en": "Clipping system DCX5FC"}, "clipper", "component", {}),
    ("touchscreen", {"de": "Deutsch/Touchscreen/Touchscreen Deutsch.docx"},
     {"de": "Touchscreen"}, "control", "component", {}),
    ("touch-panel", {"de": "Deutsch/Touchscreen/10 Bedienung Touch.docx"},
     {"de": "Bedienung Touch-Panel"}, "control", "component", {}),
    ("wartung-stoerungen", {"de": "Deutsch/Wartung_Störungen/Wartung und Störungen.docx"},
     {"de": "Wartung und Störungen"}, "service", "chapters", {"de": "Grundmaschine", "en": "Base machine"}),
    ("verbrauchsmaterial", {"de": "Deutsch/Allgemeine Anleitung/Verbrauchsgüter_Deutsch_muss angepasst werden.docx"},
     {"de": "Verbrauchsmaterial"}, "consumables", "component", {}),
    ("ce-erklaerung", {"de": "Deutsch/CE-Vordruck/CE-H10 Deutsch.docx"},
     {"de": "CE-Konformitätserklärung"}, "ce", "component", {}),
]

# options that so far only exist as error codes (no text yet)
OPTION_MODULES = {
    "option-drucker": {"de": "Option Drucker"},
    "option-waage": {"de": "Option Waage"},
    "option-sammelpacker": {"de": "Option Sammelpacker"},
    "option-beutelkontrolle": {"de": "Option Beutelkontrolle"},
}

CLIPPERS = ["clipper-dcx5s", "clipper-dcx5fc"]
DOSING = ["dosierung-servo-2k", "dosierung-elektromechanisch"]
FILM = ["folie-siegelung-2k"]
BASE = ["wartung-stoerungen"]

ASSIGN_FAULTS: dict[int, list[str]] = {}
for codes, mods in [
    ([1, 2, 6, 7, 8, 9, 11, 12, 35, 55, 56, 57, 65], CLIPPERS),
    ([4, 13, 14, 23, 26, 28, 31, 32, 67, 71, 77], FILM),
    ([3, 17, 66], DOSING),
    ([15], ["option-drucker"]),
    ([27, 99], ["option-waage"]),
    ([29, 33, 93, 94, 95, 96, 97, 98], ["option-sammelpacker"]),
    ([70, 87, 88, 90, 91, 92], ["option-beutelkontrolle"]),
]:
    for c in codes:
        ASSIGN_FAULTS[c] = mods
# everything else -> base machine (wartung-stoerungen)

# assignment of maintenance rows / spare parts by keyword in the text
ASSIGN_MAINTENANCE = [("clipper", CLIPPERS)]
ASSIGN_PARTS = [("teflon", FILM), ("temperatur sensor", FILM), ("heizpatrone", FILM)]

UNSURE = {3: "Küken – Dosierung?", 4: "Freihub – Folie?", 8: "Messer – Clipper?",
          50: "Proportional-Regelventil – Grundmaschine oder Dosierung?",
          65: "Endverdränger – Clipper?", 90: "Steckerkontrolle Beuteltisch – Beutelkontrolle?",
          93: "Puscher – Sammelpacker?", 21: "Schwenkband – Option?"}

TEXT_FIXES = [
    (re.compile(r"(siehe\s+)?Kap\.?\s*4\.4"), "Kapitel „Abschaltprozeduren“"),
    (re.compile(r"SKT[ -]H?10"), "{{Maschinentyp}}"),
]


def _cell_text(cell) -> str:
    lines = [p.text.strip() for p in cell.paragraphs]
    text = "\n".join(lines).strip("\n ")
    return re.sub(r"\n{2,}", "\n", text)


def _interval_days(text: str) -> int:
    t = text.lower()
    if "täglich" in t:
        return 1
    if "wöchentlich" in t:
        return 7
    m = re.search(r"(\d+)\s*monat", t)
    if m:
        return int(m.group(1)) * 30
    if "monatlich" in t:
        return 30
    return 9999


def _assign(text: str, rules) -> list[str]:
    low = text.lower()
    for kw, mods in rules:
        if kw in low:
            return mods
    return BASE


def _parse_service(doc, report: list[str]):
    """Extract maintenance / spare parts / fault tables -> data per module."""
    data: dict[str, dict[str, list]] = {}

    def add(mods, section, entry):
        for m in mods:
            data.setdefault(m, {}).setdefault(section, []).append(dict(entry))

    t_maint, t_parts, t_faults = doc.tables[0], doc.tables[1], doc.tables[2]
    report.append("## Wartung\n")
    for row in t_maint.rows[1:]:
        interval, task, parts = (_cell_text(c) for c in row.cells[:3])
        mods = _assign(task + " " + parts, ASSIGN_MAINTENANCE)
        add(mods, "maintenance", {"interval_days": _interval_days(interval), "interval": {"de": interval},
                                  "task": {"de": task.replace("\n", " ")}, "parts": {"de": parts}})
        report.append(f"- {interval}: {task.replace(chr(10), ' ')} / {parts.replace(chr(10), ' ')} → {', '.join(mods)}")
    report.append("\n## Verschleißteile\n")
    for row in t_parts.rows[1:]:
        interval, part_no, desc = (_cell_text(c) for c in row.cells[:3])
        mods = _assign(desc, ASSIGN_PARTS)
        add(mods, "spare_parts", {"part_no": part_no, "description": {"de": desc},
                                  "interval_days": _interval_days(interval), "interval": {"de": interval}})
        report.append(f"- {part_no} {desc} → {', '.join(mods)}")
    report.append("\n## Störungen\n")
    report.append("| Code | Meldung | Modul | prüfen |\n|---|---|---|---|")
    for row in t_faults.rows[1:]:
        head, cause, remedy = (_cell_text(c) for c in row.cells[:3])
        m = re.search(r"Fehler\s*(\d+)", head)
        if not m:
            continue
        code = int(m.group(1))
        message = re.sub(r"\s+", " ", head[m.end():].replace("\n", " ")).strip(" -/")
        message = re.sub(r"(?<=[a-zäöü])-(?=[a-zäöü])", "", message)  # "Heiz-patrone" -> "Heizpatrone"
        mods = ASSIGN_FAULTS.get(code, BASE)
        add(mods, "faults", {"code": code, "message": {"de": message},
                             "cause": {"de": cause}, "remedy": {"de": remedy}})
        report.append(f"| {code} | {message} | {', '.join(mods)} | {UNSURE.get(code, '')} |")
    return data


def _restructure_service_doc(doc) -> None:
    """Replace the three big tables by placeholders (filled per machine at build time)."""
    h2 = _heading_style_id(doc, 2)
    t_maint, t_parts, t_faults = doc.tables[0], doc.tables[1], doc.tables[2]
    t_maint._tbl.addprevious(ooxml.make_paragraph("11.0 Wartungsplan", style_id=h2))
    t_maint._tbl.addprevious(ooxml.make_paragraph(
        "Die folgende Tabelle enthält alle Wartungsarbeiten dieser Maschine. "
        "Einzelheiten stehen im Kapitel der jeweiligen Baugruppe."))
    t_maint._tbl.addprevious(ooxml.make_paragraph("{{GESAMT_WARTUNGSPLAN}}"))
    t_parts._tbl.addprevious(ooxml.make_paragraph("{{GESAMT_ERSATZTEILE}}"))
    t_faults._tbl.addprevious(ooxml.make_paragraph(
        "Die Tabelle zeigt alle Störmeldungen dieser Maschine. Ursache und Behebung finden "
        "Sie auf der angegebenen Seite im Kapitel der jeweiligen Baugruppe."))
    t_faults._tbl.addprevious(ooxml.make_paragraph("{{GESAMT_FEHLERCODES}}"))
    t_faults._tbl.addprevious(ooxml.make_paragraph("12.9 Störungen der Grundmaschine", style_id=h2))
    t_faults._tbl.addprevious(ooxml.make_paragraph("{{ABSCHNITT_STOERUNGEN}}"))
    for t in (t_maint, t_parts, t_faults):
        t._tbl.getparent().remove(t._tbl)
    for p in doc.paragraphs:
        txt = p.text.strip()
        if re.match(r"^1[12]\.1\s", txt):        # "11.1 Verschleißteile", "12.1 Maschinenbedingte Störung"
            ooxml.set_paragraph_style(p._p, h2)
    # drop empty paragraphs directly before the placeholders' tables were (cosmetic)
    body = doc.element.body
    for p in list(body.findall(qn("w:p"))):
        if not ooxml.para_text(p).strip() and not ooxml.has_graphics(p):
            nxt = p.getnext()
            if nxt is not None and nxt.tag == qn("w:p") and "{{" in ooxml.para_text(nxt):
                body.remove(p)


def _apply_text_fixes(doc) -> None:
    for pat, new in TEXT_FIXES:
        for p in ooxml.iter_paragraphs(doc.element.body):
            ooxml.replace_in_paragraph(p, func=lambda m, new=new: new, pattern=pat)


def seed_from_legacy(legacy_root: Path, library_root: Path, log=print) -> Library:
    from .importer import import_file

    template_src = legacy_root / LEGACY_MODULES[0][1]["de"]
    lib = create_library(library_root, template_src)
    report = ["# Erstimport der bestehenden Anleitungen\n",
              "Diese Zuordnung ist ein Entwurf – bitte prüfen und im Programm korrigieren "
              "(Modul doppelklicken → Tabellen).\n"]
    service_data = {}
    with tempfile.TemporaryDirectory() as tmp:
        for mid, files, title, slot, kind, component in LEGACY_MODULES:
            module = lib.new_module(mid, title, slot, kind)
            module.component = component
            if mid == "wartung-stoerungen":
                module.data_sections = True
            for lang, rel in files.items():
                src = legacy_root / rel
                doc = Document(str(src))
                if lang == "de":
                    _apply_text_fixes(doc)
                if mid == "wartung-stoerungen":
                    service_data = _parse_service(doc, report)
                    _restructure_service_doc(doc)
                fixed = Path(tmp) / f"{mid}_{lang}.docx"
                doc.save(str(fixed))
                rep = import_file(lib, fixed, module, lang, note=f"Erstimport aus {rel}")
                log(f"{mid} [{lang}]: {len(rep.headings)} Überschriften, {rep.images} Bilder")
                for i in rep.issues:
                    if i.level != "info":
                        log(f"   {i}")
            module.save_meta()
        for mid, title in OPTION_MODULES.items():
            m = lib.new_module(mid, title, "options", "component")
            m.description = "Bisher nur Störmeldungen – Beschreibung als Word-Datei ergänzen."
            m.save_meta()

    lib.reload_modules()
    for mid, sections in service_data.items():
        m = lib.modules[mid]
        for section, entries in sections.items():
            m.data[section] = entries
        m.save_data()

    report += ["\n## Nicht übernommen\n",
               "- `Englisch/Allgemeine Anleitung` und `Italienisch/…` enthalten deutschen Text "
               "(noch nicht übersetzt) und wurden nicht als Übersetzung importiert.",
               "- `Italienisch/Clipper/FC-Englisch.docx` ist englisch, identisch mit der englischen Fassung."]
    (library_root / "SEED_REPORT.md").write_text("\n".join(report) + "\n", encoding="utf-8")

    example = Machine(
        name="SKT H10 Beispiel",
        language="de",
        modules=["allgemein", "dosierung-servo-2k", "folie-siegelung-2k", "clipper-dcx5s",
                 "touchscreen", "option-waage", "wartung-stoerungen", "verbrauchsmaterial",
                 "ce-erklaerung"],
        variables={"Maschinentyp": "SKT H10", "Kunde": "", "Seriennummer": "", "Baujahr": ""},
    )
    example.save(lib.machines_dir / f"{slugify(example.name)}.yaml")
    lib.reload_modules()
    return lib
