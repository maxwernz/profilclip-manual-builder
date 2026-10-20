"""Clean up a module .docx so it can be merged into a manual.

What happens:
* headings are detected (heading styles, custom "Überschrift" styles, or bold
  numbered lines like "13.2 Drahtspezifikation")
* the level is derived from the typed number ("9.5.1" is one level below "9.5"),
  then the number is removed – Word numbers the chapters automatically later
* every heading gets the standard style "heading 1..4" and loses manual formatting
* pictures inside headings are moved below the heading
* section breaks are removed (header/footer come from the master template)
* problems that need a human are reported (hard-coded chapter references,
  machine-specific text, unknown placeholders, ...)
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn

from . import ooxml
from .library import Issue

CHAPTER_REF = re.compile(
    r"(?:Kap\.?|Kapitel|chapter|chap\.|capitolo|cap\.)\s*\d+(?:\.\d+)*", re.IGNORECASE)
TEXT_HEADING = re.compile(r"^\s*\d+(?:\.\d+)*\.?\s+[^\W\d_]")
MAX_LEVEL = 4

KNOWN_MARKERS = {
    "ABSCHNITT_WARTUNG", "ABSCHNITT_ERSATZTEILE", "ABSCHNITT_STOERUNGEN",
    "GESAMT_WARTUNGSPLAN", "GESAMT_ERSATZTEILE", "GESAMT_FEHLERCODES",
}


@dataclass
class Heading:
    level: int
    text: str
    removed_number: str | None = None


@dataclass
class NormalizeReport:
    issues: list[Issue] = field(default_factory=list)
    headings: list[Heading] = field(default_factory=list)
    images: int = 0
    words: int = 0

    def add(self, level: str, msg: str) -> None:
        self.issues.append(Issue(level, msg))

    @property
    def errors(self) -> list[Issue]:
        return [i for i in self.issues if i.level == "error"]

    def summary(self) -> str:
        lines = [f"{len(self.headings)} Überschriften, {self.images} Bilder, ~{self.words} Wörter"]
        lines += [("  " * (h.level - 1)) + f"H{h.level}  {h.text}" for h in self.headings]
        lines += [str(i) for i in self.issues]
        return "\n".join(lines)


def _style_map(doc) -> dict[str, tuple[str, int | None, bool]]:
    """style_id -> (name, heading level or None, is_bold)"""
    out = {}
    for st in doc.styles.element.findall(qn("w:style")):
        sid = st.get(qn("w:styleId"))
        name_el = st.find(qn("w:name"))
        name = name_el.get(qn("w:val")) if name_el is not None else sid
        level = None
        m = re.fullmatch(r"heading (\d)", name or "", re.IGNORECASE)
        if m:
            level = int(m.group(1))
        else:
            ol = st.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
            if ol is not None and ol.get(qn("w:val")) not in (None, "9"):
                level = int(ol.get(qn("w:val"))) + 1
        bold = st.find(qn("w:rPr") + "/" + qn("w:b")) is not None
        out[sid] = (name or sid, level, bold)
    return out


def _heading_style_id(doc, level: int) -> str:
    for st in doc.styles:
        if st.type == WD_STYLE_TYPE.PARAGRAPH and st.name.lower() == f"heading {level}":
            return st.style_id
    st = doc.styles.add_style(f"Heading {level}", WD_STYLE_TYPE.PARAGRAPH)
    return st.style_id


def _is_bold_line(p, styles) -> bool:
    runs = [r for r in p.findall(qn("w:r")) if "".join(t.text or "" for t in r.findall(qn("w:t"))).strip()]
    if not runs:
        return False
    sid = ooxml.paragraph_style_id(p)
    style_bold = styles.get(sid, ("", None, False))[2] if sid else False
    for r in runs:
        b = r.find(qn("w:rPr") + "/" + qn("w:b"))
        if b is None and not style_bold:
            return False
        if b is not None and b.get(qn("w:val")) in ("0", "false"):
            return False
    return True


def _numbered_depth(text: str) -> int | None:
    m = ooxml.NUMBER_PREFIX.match(text)
    if not m or not m.group(0).strip():
        return None
    return m.group(1).count(".") + 1


def normalize_document(doc, *, kind: str = "component", title: str = "",
                       watch_terms: list[str] | None = None,
                       known_placeholders: set[str] | None = None) -> NormalizeReport:
    rep = NormalizeReport()
    body = doc.element.body
    styles = _style_map(doc)

    # -- 1. find heading candidates among top-level paragraphs
    candidates = []   # (p, style_level or None, numbered depth or None, detected_by_text)
    for p in ooxml.body_paragraphs(body):
        text = ooxml.para_text(p).strip()
        sid = ooxml.paragraph_style_id(p)
        name, style_level, _ = styles.get(sid, ("", None, False)) if sid else ("", None, False)
        direct_ol = p.find(qn("w:pPr") + "/" + qn("w:outlineLvl"))
        if direct_ol is not None and direct_ol.get(qn("w:val")) != "9":
            style_level = int(direct_ol.get(qn("w:val"))) + 1
        is_heading_style = style_level is not None or "berschrift" in (name or "").lower() \
            or "heading" in (name or "").lower()
        if name and name.lower().startswith(("header", "footer", "kopfzeile", "fußzeile")):
            is_heading_style = False
        if is_heading_style:
            if not text:
                ooxml.set_paragraph_style(p, None)
                ooxml.remove_numbering(p)
                continue
            candidates.append((p, style_level, _numbered_depth(text), False))
        elif text and len(text) <= 90 and TEXT_HEADING.match(text) and _is_bold_line(p, styles):
            candidates.append((p, None, _numbered_depth(text), True))

    # bold "1. Step" lines are list steps, not chapters: a single-number line only
    # counts if sub-headings like "13.1" exist for it
    numbers = [ooxml.NUMBER_PREFIX.match(ooxml.para_text(p).strip()) for p, *_ in candidates]
    prefixes = {m.group(1) for m in numbers if m}
    kept = []
    for cand, m in zip(candidates, numbers):
        p, _, depth, by_text = cand
        if by_text and depth == 1 and not any(x.startswith(m.group(1) + ".") for x in prefixes):
            continue
        kept.append(cand)
    candidates = kept

    # -- 2. relative levels
    depths = [d for _, _, d, _ in candidates if d]
    base_depth = min(depths) if depths else None
    style_levels = [s for _, s, d, _ in candidates if s and not d]
    base_style = min(style_levels) if style_levels else 1
    levels = []
    prev = 0
    for p, style_level, depth, by_text in candidates:
        if depth and base_depth:
            lvl = depth - base_depth + 1
        elif style_level:
            lvl = style_level - base_style + 1 if not base_depth else max(prev, 1)
        else:
            lvl = max(prev, 1)
        levels.append(lvl)
        prev = lvl

    # -- 3. component modules form exactly one chapter
    first_is_title = False
    if kind == "component":
        ones = [i for i, lvl in enumerate(levels) if lvl == 1]
        first_block = candidates[0][0] if candidates else None
        text_before = False
        for el in body:
            if el is first_block:
                break
            if el.tag == qn("w:tbl") or (el.tag == qn("w:p") and
                                          (ooxml.para_text(el).strip() or ooxml.has_graphics(el))):
                text_before = True
                break
        first_is_title = ones == [0] and not text_before
        if not first_is_title:
            levels = [lvl + 1 for lvl in levels]
            title_p = ooxml.make_paragraph(title or "Kapitel")
            body.insert(0, title_p)
            candidates.insert(0, (title_p, 1, None, False))
            levels.insert(0, 1)
            rep.add("info", f"Kapitelüberschrift '{title}' wurde ergänzt")

    # -- 4. clamp (no skipped levels) and apply
    prev = 0
    for (p, _, depth, by_text), lvl in zip(candidates, levels):
        lvl = max(1, min(lvl, prev + 1, MAX_LEVEL))
        prev = lvl
        removed = ooxml.strip_prefix(p)
        ooxml.strip_suffix(p)
        ooxml.remove_numbering(p)
        ooxml.clear_run_formatting(p)
        pPr = p.find(qn("w:pPr"))
        if pPr is not None:
            for tag in ("w:outlineLvl", "w:pageBreakBefore"):
                el = pPr.find(qn(tag))
                if el is not None:
                    pPr.remove(el)
        ooxml.set_paragraph_style(p, _heading_style_id(doc, lvl))
        if ooxml.has_graphics(p):
            ooxml.split_graphics_out(p)
            rep.add("warning", f"Bild aus Überschrift '{ooxml.para_text(p).strip()}' "
                               f"unter die Überschrift verschoben")
        text = ooxml.para_text(p).strip()
        if by_text:
            rep.add("warning", f"'{text}' als Überschrift erkannt (fett + Nummer) – bitte prüfen")
        rep.headings.append(Heading(lvl, text, removed.strip() if removed else None))
    if not rep.headings:
        rep.add("warning", "Keine Überschriften gefunden")

    # -- 5. section breaks / headers
    for sp in body.findall(".//" + qn("w:pPr") + "/" + qn("w:sectPr")):
        orient = sp.find(qn("w:pgSz"))
        if orient is not None and orient.get(qn("w:orient")) == "landscape":
            rep.add("warning", "Querformat-Abschnitt entfernt – Seite wird hochkant")
        sp.getparent().remove(sp)
        rep.add("info", "Abschnittsumbruch entfernt (Kopf-/Fußzeile kommt aus der Vorlage)")

    # -- 6. content checks
    all_text = []
    for p in ooxml.iter_paragraphs(body):
        all_text.append(ooxml.para_text(p))
    joined = "\n".join(all_text)
    rep.words = len(joined.split())
    rep.images = sum(1 for r in doc.part.rels.values() if r.reltype.endswith("/image"))
    for m in sorted(set(CHAPTER_REF.findall(joined))):
        rep.add("warning", f"Feste Kapitelangabe '{m}' – die Nummer kann sich ändern")
    for term in watch_terms or []:
        if term.lower() in joined.lower():
            rep.add("warning", f"Maschinenspezifischer Text '{term}' – besser eine Variable "
                               f"wie {{{{Maschinentyp}}}} verwenden")
    known = (known_placeholders or set()) | KNOWN_MARKERS
    for name in sorted(set(ooxml.PLACEHOLDER.findall(joined))):
        if name not in known and not name.startswith("L:"):
            rep.add("warning", f"Unbekannter Platzhalter {{{{{name}}}}}")
    return rep


def normalize_file(src: Path, dst: Path | None = None, **kw) -> NormalizeReport:
    doc = Document(str(src))
    rep = normalize_document(doc, **kw)
    if dst is not None:
        doc.save(str(dst))
    return rep
