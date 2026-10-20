"""Assemble a manual: master template + modules + generated tables -> DOCX/PDF."""

from __future__ import annotations

import datetime as _dt
import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxcompose.composer import Composer

from . import ooxml
from .engines import finalize
from .library import DATA_SECTIONS, Issue, Library, Machine, Module, tr
from .normalize import normalize_document
from .tables import make_table, pageref
from .template import heading_style_ids

Log = Callable[[str], None]

MARKER = re.compile(r"\{\{\s*(ABSCHNITT_WARTUNG|ABSCHNITT_ERSATZTEILE|ABSCHNITT_STOERUNGEN|"
                    r"AUTO_ABSCHNITTE|GESAMT_WARTUNGSPLAN|GESAMT_ERSATZTEILE|GESAMT_FEHLERCODES)"
                    r"(?::([A-Za-z0-9_\-]+))?\s*\}\}")
OWN_MARKER = {"ABSCHNITT_WARTUNG": "maintenance", "ABSCHNITT_ERSATZTEILE": "spare_parts",
              "ABSCHNITT_STOERUNGEN": "faults"}
SECTION_LABEL = {"maintenance": "maintenance", "spare_parts": "spare_parts", "faults": "faults"}


class BuildError(RuntimeError):
    def __init__(self, issues: list[Issue]):
        super().__init__("\n".join(str(i) for i in issues))
        self.issues = issues


@dataclass
class BuildResult:
    docx: Path
    pdf: Path | None
    engine: str
    issues: list[Issue] = field(default_factory=list)


def output_basename(machine: Machine) -> str:
    safe = re.sub(r"[^\w\-]+", "_", machine.name).strip("_") or "Anleitung"
    return f"{safe}_{machine.language}"


# ---------------------------------------------------------------- data rendering

class _Renderer:
    def __init__(self, lib: Library, machine: Machine, doc):
        self.lib = lib
        self.lang = machine.language
        self.doc = doc
        self.modules = [lib.modules[m] for m in machine.modules if m in lib.modules]
        self.heading_ids = heading_style_ids(doc)
        self.fault_bookmarks: dict[str, str] = {}
        self.untranslated: set[str] = set()

    def L(self, key: str) -> str:
        return self.lib.label(key, self.lang)

    def t(self, value, module: Module) -> str:
        text, fallback = tr(value, self.lang, self.lib.default_language)
        if fallback:
            self.untranslated.add(module.label(self.lang))
        return text

    # one module's own table
    def own_table(self, module: Module, section: str):
        e = module.entries(section)
        if section == "maintenance":
            e.sort(key=lambda x: x.get("interval_days") or 0)
            return make_table([self.L("interval"), self.L("task"), self.L("parts")],
                              [[self.t(x.get("interval"), module), self.t(x.get("task"), module),
                                self.t(x.get("parts"), module)] for x in e],
                              [3.5, 5.5, 7.0])
        if section == "spare_parts":
            return make_table([self.L("part_no"), self.L("description"), self.L("interval")],
                              [[str(x.get("part_no", "")), self.t(x.get("description"), module),
                                self.t(x.get("interval"), module)] for x in e],
                              [3.0, 8.5, 4.5])
        e.sort(key=lambda x: _code_key(x.get("code")))
        return make_table([self.L("code"), self.L("cause"), self.L("remedy")],
                          [[f"{x.get('code', '')}\n{self.t(x.get('message'), module)}",
                            self.t(x.get("cause"), module), self.t(x.get("remedy"), module)]
                           for x in e],
                          [4.0, 5.5, 6.5])

    def own_sections(self, module: Module, with_heading: bool, only: str | None = None) -> list:
        out = []
        for section in DATA_SECTIONS:
            if only and section != only:
                continue
            if not module.entries(section):
                continue
            anchor = ooxml.make_paragraph(self.L(SECTION_LABEL[section]) if with_heading else "",
                                          style_id=self.heading_ids.get(2) if with_heading else None)
            if section == "faults":
                bm = ooxml.bookmark_name("flt", module.id)
                ooxml.add_bookmark(anchor, bm)
                self.fault_bookmarks[module.id] = bm
            if with_heading or section == "faults":
                out.append(anchor)
            out.append(self.own_table(module, section))
            out.append(ooxml.make_paragraph(""))
        return out

    # overviews across all modules of the machine
    def overview(self, kind: str) -> list:
        rows = []
        if kind == "GESAMT_WARTUNGSPLAN":
            for m in self.modules:
                for x in m.entries("maintenance"):
                    rows.append((x.get("interval_days") or 0, [self.t(x.get("interval"), m), m.component_label(self.lang),
                                                                self.t(x.get("task"), m), self.t(x.get("parts"), m)]))
            rows.sort(key=lambda r: r[0])
            tbl = make_table([self.L("interval"), self.L("module"), self.L("task"), self.L("parts")],
                             [r for _, r in rows], [3.0, 3.5, 4.0, 5.5])
        elif kind == "GESAMT_ERSATZTEILE":
            for m in self.modules:
                for x in m.entries("spare_parts"):
                    rows.append((str(x.get("part_no", "")), [str(x.get("part_no", "")),
                                                              self.t(x.get("description"), m),
                                                              self.t(x.get("interval"), m), m.component_label(self.lang)]))
            rows.sort(key=lambda r: r[0])
            tbl = make_table([self.L("part_no"), self.L("description"), self.L("interval"), self.L("module")],
                             [r for _, r in rows], [2.8, 6.0, 3.7, 3.5])
        else:
            for m in self.modules:
                for x in m.entries("faults"):
                    rows.append((_code_key(x.get("code")), [str(x.get("code", "")), self.t(x.get("message"), m),
                                                             m.component_label(self.lang),
                                                             pageref(self.fault_bookmarks.get(m.id))]))
            rows.sort(key=lambda r: r[0])
            tbl = make_table([self.L("no"), self.L("message"), self.L("module"), self.L("page")],
                             [r for _, r in rows], [1.5, 8.0, 4.5, 2.0])
        return [tbl, ooxml.make_paragraph("")]


def _code_key(code) -> tuple:
    s = str(code or "")
    m = re.match(r"\d+", s)
    return (int(m.group()) if m else 10 ** 6, s)


# ---------------------------------------------------------------- build

def _tag_module_markers(doc, module: Module) -> bool:
    """{{ABSCHNITT_X}} -> {{ABSCHNITT_X:<module id>}}; returns True if any was found."""
    found = [False]

    def repl(m):
        if m.group(1) in OWN_MARKER and not m.group(2):
            found[0] = True
            return "{{%s:%s}}" % (m.group(1), module.id)
        return None

    for p in ooxml.iter_paragraphs(doc.element.body):
        ooxml.replace_in_paragraph(p, func=repl, pattern=MARKER)
    return found[0]


def _body_append(doc, el) -> None:
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        sect.addprevious(el)
    else:
        body.append(el)


def _empty_doc_like(template: Path):
    doc = Document(str(template))
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    return doc


def _replace_markers(doc, renderer: _Renderer, issues: list[Issue]) -> None:
    body = doc.element.body
    # pass 1: module sections (creates the bookmarks the overview refers to)
    # pass 2: overviews
    for pass_no in (1, 2):
        for p in list(body.iter(qn("w:p"))):
            text = ooxml.para_text(p).strip()
            m = MARKER.fullmatch(text) if text else None
            if not m:
                continue
            kind, mid = m.group(1), m.group(2)
            is_overview = kind.startswith("GESAMT_")
            if (pass_no == 1) == is_overview:
                continue
            if is_overview:
                new = renderer.overview(kind)
            else:
                module = renderer.lib.modules.get(mid or "")
                if module is None:
                    issues.append(Issue("warning", f"Platzhalter {text} ohne Modul"))
                    new = []
                elif kind == "AUTO_ABSCHNITTE":
                    new = renderer.own_sections(module, with_heading=True)
                else:
                    new = renderer.own_sections(module, with_heading=False, only=OWN_MARKER[kind])
            for el in new:
                p.addprevious(el)
            p.getparent().remove(p)


def _replace_placeholders(doc, values: dict[str, str], lib: Library, lang: str,
                          issues: list[Issue]) -> None:
    missing: set[str] = set()

    def repl(m):
        key = m.group(1)
        if key.startswith("L:"):
            return lib.label(key[2:], lang)
        if key in values:
            return values[key]
        if MARKER.fullmatch(m.group(0)):
            return None
        missing.add(key)
        return None

    roots = [doc.element.body]
    for section in doc.sections:
        for hf in (section.header, section.footer, section.first_page_header,
                   section.first_page_footer, section.even_page_header, section.even_page_footer):
            if hf is not None and not hf.is_linked_to_previous:
                roots.append(hf._element)
    for root in roots:
        for p in ooxml.iter_paragraphs(root):
            ooxml.replace_in_paragraph(p, func=repl)
    for key in sorted(missing):
        issues.append(Issue("warning", f"Platzhalter {{{{{key}}}}} hat keinen Wert"))


def _set_update_fields(doc) -> None:
    settings = doc.settings.element
    el = settings.find(qn("w:updateFields"))
    if el is None:
        el = OxmlElement("w:updateFields")
        settings.append(el)
    el.set(qn("w:val"), "true")


def compose(lib: Library, machine: Machine, out_docx: Path, log: Log = print) -> list[Issue]:
    """Write the merged .docx (fields not yet updated). Returns warnings."""
    issues = lib.validate(machine)
    errors = [i for i in issues if i.level == "error"]
    if errors:
        raise BuildError(errors)
    lang = machine.language
    template = lib.template_path(lang)
    master = Document(str(template))
    composer = Composer(master)
    watch = list(lib.config.get("watch_terms") or [])
    known = {v["key"] for v in lib.variables} | {"Ausgabe"}

    for mid in machine.modules:
        module = lib.modules[mid]
        title = module.label(lang)
        log(f"+ {title}")
        if module.has_text():
            src = module.docx_path(lang)
            if not src.exists():
                src = module.docx_path(lib.default_language)
                if not src.exists():
                    src = module.docx_path(module.languages()[0])
            doc = Document(str(src))
            rep = normalize_document(doc, kind=module.kind, title=title, watch_terms=watch,
                                     known_placeholders=known)
            for i in rep.issues:
                if i.level != "info" and "als Überschrift erkannt" not in i.message:
                    issues.append(Issue(i.level, f"{title}: {i.message}"))
        else:
            doc = _empty_doc_like(template)
            ids = heading_style_ids(doc)
            _body_append(doc, ooxml.make_paragraph(title, style_id=ids.get(1)))
        has_own = _tag_module_markers(doc, module)
        if module.data_sections and module.has_data() and not has_own:
            _body_append(doc, ooxml.make_paragraph("{{AUTO_ABSCHNITTE:%s}}" % module.id))
        composer.append(doc)

    final = composer.doc
    renderer = _Renderer(lib, machine, final)
    _replace_markers(final, renderer, issues)
    if renderer.untranslated:
        issues.append(Issue("warning", "Daten ohne Übersetzung (deutscher Text verwendet): "
                                       + ", ".join(sorted(renderer.untranslated))))
    codes: dict[str, list[str]] = {}
    for m in renderer.modules:
        for x in m.entries("faults"):
            codes.setdefault(str(x.get("code")), []).append(m.label(lang))
    for code, mods in codes.items():
        if len(mods) > 1:
            issues.append(Issue("warning", f"Fehler {code} kommt mehrfach vor: {', '.join(mods)}"))

    values = lib.default_variables(lang)
    values.update({k: v for k, v in machine.variables.items() if v != ""})
    if not values.get("Ausgabe"):
        values["Ausgabe"] = _dt.date.today().strftime("%d.%m.%Y")
    _replace_placeholders(final, values, lib, lang, issues)
    _set_update_fields(final)
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    final.save(str(out_docx))
    return issues


def build_manual(lib: Library, machine: Machine, out_dir: Path, *, pdf: bool = True,
                 engine: str = "auto", log: Log = print) -> BuildResult:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    base = output_basename(machine)
    out_docx = out_dir / f"{base}.docx"
    out_pdf = out_dir / f"{base}.pdf" if pdf else None
    with tempfile.TemporaryDirectory(prefix="manual-builder-") as tmp:
        raw = Path(tmp) / f"{base}.docx"
        issues = compose(lib, machine, raw, log)
        log("Inhaltsverzeichnis und Seitenzahlen aktualisieren …")
        used = finalize(raw, out_docx, out_pdf, engine=engine, log=log)
    if out_pdf and not out_pdf.exists():
        out_pdf = None
    log(f"Fertig: {out_docx.name}" + (f" + {out_pdf.name}" if out_pdf else "") + f" (Engine: {used})")
    return BuildResult(out_docx, out_pdf, used, issues)
