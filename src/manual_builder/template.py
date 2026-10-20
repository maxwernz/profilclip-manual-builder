"""Create / prepare the master template (cover, table of contents, numbered headings).

The template is an ordinary .docx and can be edited in Word afterwards (logo,
fonts, cover layout). Only three things must stay:
  * the heading styles "Überschrift 1-4" (heading 1-4) – they carry the numbering
  * the table-of-contents field
  * the {{...}} placeholders you want filled in
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import ooxml
from .tables import make_table

# text in the existing header/footer that becomes a placeholder
HEADER_REPLACEMENTS = [
    ("Originalbetriebsanleitung", "{{Dokumentart}}"),
    ("Schlauchbeutel-Abfüllautomat", "{{Maschinenbezeichnung}}"),
    ("SKT H10", "{{Maschinentyp}}"),
    ("03810000", "{{Dokumentnummer}}"),
    ("Gedruckt", "{{L:printed}}"),
]

HEADING_SIZES = {1: 16, 2: 14, 3: 12, 4: 11}


def heading_style_ids(doc) -> dict[int, str]:
    out = {}
    for st in doc.styles:
        if st.type == WD_STYLE_TYPE.PARAGRAPH:
            m = re.fullmatch(r"heading (\d)", st.name.lower())
            if m:
                out[int(m.group(1))] = st.style_id
    return out


def _ensure_heading_styles(doc) -> dict[int, str]:
    ids = heading_style_ids(doc)
    for lvl in range(1, 5):
        if lvl not in ids:
            st = doc.styles.add_style(f"Heading {lvl}", WD_STYLE_TYPE.PARAGRAPH)
            ids[lvl] = st.style_id
    return ids


def _numbering_part(doc):
    try:
        return doc.part.numbering_part
    except (KeyError, NotImplementedError):
        raise RuntimeError("Vorlage hat keine Nummerierungsdefinition (numbering.xml)")


def _add_heading_numbering(doc, style_ids: dict[int, str]) -> int:
    numbering = _numbering_part(doc).element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    abs_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    a = OxmlElement("w:abstractNum")
    a.set(qn("w:abstractNumId"), str(abs_id))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "multilevel")
    a.append(mlt)
    for ilvl in range(9):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", "decimal")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        if ilvl + 1 in style_ids:
            e = OxmlElement("w:pStyle")
            e.set(qn("w:val"), style_ids[ilvl + 1])
            lvl.append(e)
        txt = OxmlElement("w:lvlText")
        txt.set(qn("w:val"), ".".join(f"%{i + 1}" for i in range(ilvl + 1)))
        lvl.append(txt)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        pPr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        width = (680, 851, 1021, 1191)[min(ilvl, 3)]
        ind.set(qn("w:left"), str(width))
        ind.set(qn("w:hanging"), str(width))
        pPr.append(ind)
        lvl.append(pPr)
        a.append(lvl)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(a)
    else:
        numbering.append(a)
    n = OxmlElement("w:num")
    n.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abs_id))
    n.append(ref)
    numbering.append(n)
    return num_id


def _style_headings(doc, style_ids: dict[int, str], num_id: int, font: str) -> None:
    for lvl, sid in style_ids.items():
        if lvl > 4:
            continue
        st = doc.styles.get_by_id(sid, WD_STYLE_TYPE.PARAGRAPH)
        st.font.name = font
        st.font.size = Pt(HEADING_SIZES[lvl])
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor(0, 0, 0)
        rPr = st.element.get_or_add_rPr()
        for tag in ("w:rFonts",):
            rf = rPr.find(qn(tag))
            if rf is not None:
                for att in list(rf.attrib):
                    if "Theme" in att:
                        del rf.attrib[att]
                for att in ("ascii", "hAnsi", "cs", "eastAsia"):
                    rf.set(qn(f"w:{att}"), font)
        pf = st.paragraph_format
        pf.space_before = Pt(18 if lvl == 1 else 12)
        pf.space_after = Pt(6)
        pf.keep_with_next = True
        pf.page_break_before = lvl == 1
        pPr = st.element.get_or_add_pPr()
        old = pPr.find(qn("w:numPr"))
        if old is not None:
            pPr.remove(old)
        numPr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(lvl - 1))
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        numPr.append(ilvl)
        numPr.append(nid)
        pPr.insert(0, numPr)
        ol = pPr.find(qn("w:outlineLvl"))
        if ol is None:
            ol = OxmlElement("w:outlineLvl")
            pPr.append(ol)
        ol.set(qn("w:val"), str(lvl - 1))


def _clear_body(doc) -> None:
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def _append(doc, el) -> None:
    body = doc.element.body
    sect = body.find(qn("w:sectPr"))
    if sect is not None:
        sect.addprevious(el)
    else:
        body.append(el)


def _replace_literal(part_element, pairs) -> None:
    for old, new in pairs:
        pat = re.compile(re.escape(old))
        for p in ooxml.iter_paragraphs(part_element):
            ooxml.replace_in_paragraph(p, func=lambda m, new=new: new, pattern=pat)


def create_master_template(source_docx: Path, out: Path, font: str = "Arial") -> Path:
    """Build master.docx from an existing manual (keeps page setup, header, footer, logo)."""
    doc = Document(str(source_docx))
    _clear_body(doc)
    for section in doc.sections:
        for hf in (section.header, section.footer, section.first_page_header,
                   section.first_page_footer, section.even_page_header, section.even_page_footer):
            if hf is not None and not hf.is_linked_to_previous:
                _replace_literal(hf._element, HEADER_REPLACEMENTS)

    style_ids = _ensure_heading_styles(doc)
    num_id = _add_heading_numbering(doc, style_ids)
    _style_headings(doc, style_ids, num_id, font)

    # cover page
    for _ in range(6):
        _append(doc, ooxml.make_paragraph(""))
    _append(doc, ooxml.make_paragraph("{{Dokumentart}}", size_half_points=28, align="center"))
    _append(doc, ooxml.make_paragraph(""))
    _append(doc, ooxml.make_paragraph("{{Maschinenbezeichnung}}", bold=True, size_half_points=40,
                                      align="center"))
    _append(doc, ooxml.make_paragraph("{{Maschinentyp}}", bold=True, size_half_points=56,
                                      align="center"))
    for _ in range(8):
        _append(doc, ooxml.make_paragraph(""))
    rows = [["{{L:customer}}", "{{Kunde}}"],
            ["{{L:serial}}", "{{Seriennummer}}"],
            ["{{L:year}}", "{{Baujahr}}"],
            ["{{L:document_no}}", "{{Dokumentnummer}}"],
            ["{{L:issue}}", "{{Ausgabe}}"]]
    _append(doc, make_table(["{{L:machine_data}}", ""], rows, [5.5, 9.5], font_size_pt=11))

    # table of contents on its own page
    _append(doc, ooxml.make_paragraph("{{L:contents}}", bold=True, size_half_points=32,
                                      page_break_before=True))
    toc = ooxml.make_paragraph("")
    ooxml.add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "{{L:toc_hint}}")
    _append(doc, toc)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
